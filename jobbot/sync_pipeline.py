"""Sync the resume pipeline and the application tracker between OneDrive and GitHub.

OneDrive's Applications 2026 stays the place the user edits. GitHub's runners
cannot see it, so this:
  1. pulls jobs he approved on his phone (data/approvals.jsonl) and appends any
     not yet in Tracker.xlsx, so the tracker stays his single record;
  2. copies what the cloud build needs into pipeline/;
  3. exports every application - Tracker.xlsx rows and <Company>/<Role>/JD.md
     folders - to pipeline/prior_applications.json, so the cloud never alerts
     on, lists, or builds a resume for a job he already applied to.

    python -m jobbot.sync_pipeline          # all of the above + push
    python -m jobbot.sync_pipeline --no-push

Rerun after editing profile.yaml, PIPELINE.md, resume.cls or verify_resume.py,
and after applying to anything manually.
"""
import argparse, json, os, re, shutil, subprocess, sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path(os.environ.get("JOBBOT_APPS") or
              Path.home() / "OneDrive" / "Desktop" / "Resumes" / "Applications 2026")
DEST = ROOT / "pipeline"
FILES = ["profile.yaml", "PIPELINE.md", "resume.cls", "scripts/verify_resume.py"]

# Job ids quoted in a JD.md header: long digit runs (Greenhouse) and UUIDs (Lever/Ashby).
ID_RE = re.compile(r"\b(\d{6,}|[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\b", re.I)


TRACKER_SHEET = "Internships Applied"


def _key(company, title):
    norm = lambda s: re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()
    return norm(company), norm(title)


def tracker_rows(src):
    import openpyxl
    wb = openpyxl.load_workbook(src / "Tracker.xlsx", read_only=True)
    rows = [r for r in wb[TRACKER_SHEET].iter_rows(min_row=2, values_only=True) if r and r[0]]
    wb.close()
    return rows


def import_approvals(src):
    """Append phone approvals missing from Tracker.xlsx. Returns rows added."""
    path = ROOT / "data" / "approvals.jsonl"
    if not path.exists():
        return 0
    approvals = [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]
    have = {_key(r[0], r[1]) for r in tracker_rows(src)}
    todo = [a for a in approvals if _key(a.get("company"), a.get("title")) not in have]
    if not todo:
        return 0
    import openpyxl
    tracker = src / "Tracker.xlsx"
    shutil.copy2(tracker, tracker.with_name("Tracker.xlsx.bak-jobbot"))
    wb = openpyxl.load_workbook(tracker)
    ws = wb[TRACKER_SHEET]
    row = ws.max_row
    while row > 0 and ws.cell(row=row, column=1).value is None:
        row -= 1
    for a in todo:
        row += 1
        when = datetime.fromisoformat(a["approved_at"]).replace(tzinfo=None, hour=0, minute=0, second=0)
        values = [a.get("company"), a.get("title"), when, "None", "No Response",
                  f"Approved in jobbot (issue #{a.get('issue')}). Resume: {a.get('pdf')}. "
                  f"Apply: {a.get('apply_url')}"]
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v)
    wb.save(tracker)
    return len(todo)


def prior_applications(src):
    out = []
    for jd in sorted(src.glob("*/*/JD.md")):
        text = jd.read_text(encoding="utf-8", errors="replace")[:6000]
        out.append({"company": jd.parent.parent.name, "role": jd.parent.name,
                    "ids": sorted(set(ID_RE.findall(text))), "source": "folder"})
    seen = {_key(p["company"], p["role"]) for p in out}
    for r in tracker_rows(src):
        if _key(r[0], r[1]) not in seen:
            out.append({"company": str(r[0]), "role": str(r[1] or ""), "ids": [], "source": "tracker"})
    return out


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-push", action="store_true")
    a = ap.parse_args(argv)
    if not SOURCE.exists():
        print(f"ERROR: {SOURCE} not found", file=sys.stderr)
        return 1

    git = lambda *args: subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)
    git("fetch", "-q", "origin", "main")
    shown = git("show", "origin/main:data/approvals.jsonl")
    if shown.returncode == 0:
        (ROOT / "data" / "approvals.jsonl").write_text(shown.stdout, encoding="utf-8")
    try:
        added = import_approvals(SOURCE)
        print(f"Tracker.xlsx: {added} phone approval(s) added")
    except PermissionError:
        print("ERROR: Tracker.xlsx is open in Excel - close it and rerun", file=sys.stderr)
        return 1

    for rel in FILES:
        (DEST / rel).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(SOURCE / rel, DEST / rel)
    prior = prior_applications(SOURCE)
    (DEST / "prior_applications.json").write_text(json.dumps(prior, indent=1), encoding="utf-8")
    print(f"synced {len(FILES)} files and {len(prior)} prior applications "
          f"({sum(p['source'] == 'tracker' for p in prior)} from Tracker.xlsx) into {DEST}")
    if a.no_push:
        return 0
    from . import gitsync
    return gitsync.push("sync resume pipeline and tracker from OneDrive", ["pipeline"])


if __name__ == "__main__":
    sys.exit(main())
